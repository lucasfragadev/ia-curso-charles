"""
Script de Auditoria (Conferencia-Modelo01)

Cruza a planilha de controle manual com a base de dados oficial
gerada a partir do SIGTAP para um determinado CBO.

Saida:
    Uma planilha na pasta `auditoria/YYYY-MM-DD/` com timestamp unico
    no nome para evitar sobrescrita entre execucoes do mesmo dia.
"""
import argparse
import os
from datetime import datetime
import pandas as pd
from openpyxl import load_workbook

from src.config import CBO_NEUROCIRURGIAO, PASTA_DADOS


def _detectar_cabecalho(planilha_path: str) -> int:
    df_temp = pd.read_excel(planilha_path, header=None, nrows=25, engine='openpyxl')
    best_row = 0
    max_score = 0
    for idx, row in df_temp.iterrows():
        row_str = ' '.join([str(val).upper() for val in row.values if pd.notna(val)])
        score = 0
        if 'SIGTAP' in row_str: score += 5
        if 'VALOR' in row_str: score += 2
        if 'PACIENTE' in row_str: score += 2
        if 'CÓD' in row_str or 'COD' in row_str: score += 2
        if score > max_score:
            max_score = score
            best_row = idx
    return best_row


def executar(planilha_path: str, cbo: str = CBO_NEUROCIRURGIAO, mapeamento: dict = None, linha_cabecalho: int = None) -> dict:
    """
    Executa a auditoria Modelo 01 de forma programatica.

    Args:
        planilha_path : Caminho completo para a planilha de controle.
        cbo           : Codigo CBO alvo.
        mapeamento    : Dicionario com os nomes reais das colunas na planilha.
        linha_cabecalho: Indice da linha de cabecalho. Se None, detecta automaticamente.
    """
    base_verdade_path = os.path.join(PASTA_DADOS, f"exclusivo - procedimentos_cbo_{cbo}_com_valores.xlsx")

    if not os.path.exists(planilha_path):
        return {"sucesso": False, "erro": f"Planilha '{planilha_path}' nao encontrada."}
    if not os.path.exists(base_verdade_path):
        return {"sucesso": False, "erro": f"Base da verdade '{base_verdade_path}' nao encontrada. Verifique se a pasta 'data/' esta no mesmo local que o programa."}

    # Nome de arquivo unico com timestamp para nao sobrescrever execucoes anteriores
    agora = datetime.now()
    pasta_auditoria = os.path.join("auditoria", agora.strftime("%Y-%m-%d"))
    os.makedirs(pasta_auditoria, exist_ok=True)
    arquivo_saida = os.path.join(
        pasta_auditoria,
        f"auditoria_modelo01_CBO_{cbo}_{agora.strftime('%H-%M-%S')}.xlsx"
    )

    # 1. Ler a Base da Verdade (SIGTAP)
    df_sigtap = pd.read_excel(base_verdade_path, dtype={"CODIGO": str}, engine='openpyxl')
    df_sigtap["CODIGO"] = df_sigtap["CODIGO"].astype(str).str.strip().str.zfill(10)
    dict_valores_sp = dict(zip(df_sigtap["CODIGO"], df_sigtap["valorSP"]))

    # 2. Ler a Planilha (detectando cabeçalho automaticamente)
    if linha_cabecalho is None:
        linha_cabecalho = _detectar_cabecalho(planilha_path)
    df = pd.read_excel(planilha_path, header=linha_cabecalho, dtype=str, engine='openpyxl')

    # 3. Determinar colunas — mapeamento manual tem prioridade sobre auto-deteccao
    col_codigo = None
    col_valor = None

    if mapeamento and mapeamento.get("codigo") and mapeamento.get("valor"):
        col_codigo = mapeamento["codigo"]
        col_valor = mapeamento["valor"]
    else:
        # Auto-deteccao por nome
        for col in df.columns:
            col_upper = str(col).upper()
            if "SIGTAP" in col_upper and ("COD" in col_upper or "CÓD" in col_upper or "C" in col_upper):
                col_codigo = col_codigo or col
            if "VALOR" in col_upper and "SP" in col_upper:
                col_valor = col_valor or col

    if not col_codigo or not col_valor:
        colunas_disponiveis = [str(c) for c in df.columns.tolist() if not str(c).startswith("Unnamed")]
        return {
            "sucesso": False,
            "erro": (
                "Nao foi possivel identificar as colunas automaticamente. "
                f"Colunas encontradas na planilha: {colunas_disponiveis}"
            )
        }

    # Coluna de nome do paciente: segunda coluna disponivel por padrao
    col_nome = df.columns[1] if len(df.columns) > 1 else df.columns[0]

    # 4. Processar Auditoria
    resultados = []

    for _, row in df.iterrows():
        cod_sigtap_raw = str(row.get(col_codigo, "")).strip()

        # Ignorar linhas vazias ou de total
        if cod_sigtap_raw.upper() in ("NAN", "NAT", "<NA>", "NONE", ""):
            continue
        if "TOTAL" in str(row.get("Descritivo", row.iloc[0] if len(row) > 0 else "")).upper():
            continue

        # Normalizar codigo SIGTAP
        if cod_sigtap_raw.endswith(".0"):
            cod_sigtap_raw = cod_sigtap_raw[:-2]
        if not cod_sigtap_raw.replace(".", "").isdigit():
            continue

        cod_sigtap = cod_sigtap_raw.zfill(10)

        try:
            valor_planilha = float(str(row.get(col_valor, "0")).replace(",", ".").strip())
        except (ValueError, TypeError):
            valor_planilha = 0.0

        if cod_sigtap not in dict_valores_sp:
            status = "Procedimento nao previsto para o CBO trabalhado"
            valor_datasus = None
            diferenca = None
        else:
            valor_datasus = float(dict_valores_sp[cod_sigtap])
            diferenca = valor_datasus - valor_planilha
            status = "Valores corretos" if abs(diferenca) < 0.05 else "Valores divergentes"

        resultados.append({
            "Paciente":           row.get(col_nome, ""),
            "Código Sigtap":      cod_sigtap,
            "Valor SP (Planilha)": valor_planilha,
            "Valor SP (DATASUS)": valor_datasus if valor_datasus is not None else "N/A",
            "Status Auditoria":   status,
            "Diferença (R$)":     diferenca if diferenca is not None else "N/A",
        })

    # 5. Salvar Relatorio
    df_resultado = pd.DataFrame(resultados)
    df_resultado.to_excel(arquivo_saida, index=False)

    wb = load_workbook(arquivo_saida)
    ws = wb.active
    fmt = 'R$ #,##0.00'
    idx_pl = df_resultado.columns.get_loc("Valor SP (Planilha)") + 1
    idx_dt = df_resultado.columns.get_loc("Valor SP (DATASUS)") + 1
    idx_df = df_resultado.columns.get_loc("Diferença (R$)") + 1

    for r in range(2, len(df_resultado) + 2):
        ws.cell(r, idx_pl).number_format = fmt
        c_dt = ws.cell(r, idx_dt)
        if c_dt.value != "N/A":
            c_dt.number_format = fmt
        c_df = ws.cell(r, idx_df)
        if c_df.value != "N/A":
            c_df.number_format = fmt

    wb.save(arquivo_saida)

    total        = len(df_resultado)
    corretos     = int((df_resultado["Status Auditoria"] == "Valores corretos").sum())
    divergentes  = int((df_resultado["Status Auditoria"] == "Valores divergentes").sum())
    nao_previstos = int(df_resultado["Status Auditoria"].str.contains("nao previsto", case=False, na=False).sum())

    return {
        "sucesso": True,
        "total": total,
        "corretos": corretos,
        "divergentes": divergentes,
        "nao_previstos": nao_previstos,
        "perc_principal_ok": "N/A",
        "perc_auxiliar_ok": "N/A",
        "arquivo_saida": os.path.abspath(arquivo_saida),
    }


def main():
    parser = argparse.ArgumentParser(description="Auditoria de valores SIGTAP - Modelo 01")
    parser.add_argument("--planilha",    default="docs/1- JULHO  NEUROCIRURGIA 2026 LETICIA.xlsx")
    parser.add_argument("--cbo",         default=CBO_NEUROCIRURGIAO)
    parser.add_argument("--col_codigo",  default=None, help="Nome exato da coluna com o Codigo SIGTAP")
    parser.add_argument("--col_valor",   default=None, help="Nome exato da coluna com o Valor SP")
    args = parser.parse_args()

    mapeamento = None
    if args.col_codigo and args.col_valor:
        mapeamento = {"codigo": args.col_codigo, "valor": args.col_valor}

    resultado = executar(args.planilha, args.cbo, mapeamento)

    if not resultado["sucesso"]:
        print(f"[!] Erro: {resultado['erro']}")
        return

    print("\n" + "=" * 60)
    print("Auditoria Modelo 01 concluida!")
    print(f"  Total          : {resultado['total']}")
    print(f"  Corretos       : {resultado['corretos']}")
    print(f"  Divergentes    : {resultado['divergentes']}")
    print(f"  Fora do CBO    : {resultado['nao_previstos']}")
    print(f"  Relatorio      : {resultado['arquivo_saida']}")
    print("=" * 60)


if __name__ == "__main__":
    main()
