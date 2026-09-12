"""
Script de Auditoria (Conferencia-Modelo02)

Cruza a planilha de produtividade (formato "PLENA") com a base de dados
oficial gerada a partir do SIGTAP para um determinado CBO.

Validacoes realizadas por linha:
    1. Codigo SIGTAP: procedimento previsto para o CBO?
    2. Valor do Procedimento: confere com o valorSP do SIGTAP?
    3. Valor Cirurgiao Principal: equivale a 75% do valor do procedimento?
    4. Valor Medico Auxiliar: equivale a 25% do valor do procedimento?

Saida:
    Uma planilha na pasta `auditoria/YYYY-MM-DD/` com timestamp unico
    no nome para evitar sobrescrita entre execucoes do mesmo dia.
"""
import argparse
import os
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font

from src.config import CBO_NEUROCIRURGIAO, PASTA_DADOS

# =====================================================================
# Constantes de validacao
# =====================================================================
PERCENTUAL_PRINCIPAL = 0.75
PERCENTUAL_AUXILIAR  = 0.25
TOLERANCIA_CENTAVOS  = 0.05


# =====================================================================
# Funcoes auxiliares
# =====================================================================

def _detectar_cabecalho(planilha_path: str) -> int:
    df_temp = pd.read_excel(planilha_path, header=None, nrows=25, engine='openpyxl')
    best_row = 0
    max_score = 0
    for idx, row in df_temp.iterrows():
        row_str = ' '.join([str(val).upper() for val in row.values if pd.notna(val)])
        score = 0
        if 'SIGTAP' in row_str: score += 5
        if 'VALOR' in row_str: score += 2
        if 'PACIENTE' in row_str: score += 2
        if 'CÓD' in row_str or 'COD' in row_str: score += 2
        if score > max_score:
            max_score = score
            best_row = idx
    return best_row


def _carregar_base_sigtap(cbo: str) -> dict:
    base_path = os.path.join(PASTA_DADOS, f"exclusivo - procedimentos_cbo_{cbo}_com_valores.xlsx")
    if not os.path.exists(base_path):
        raise FileNotFoundError(
            f"Base da verdade nao encontrada: '{base_path}'. "
            "Verifique se a pasta 'data/' esta no mesmo local que o programa."
        )
    df = pd.read_excel(base_path, dtype={"CODIGO": str}, engine='openpyxl')
    df["CODIGO"] = df["CODIGO"].astype(str).str.strip().str.zfill(10)
    return dict(zip(df["CODIGO"], df["valorSP"]))


def _carregar_planilha_plena(caminho: str, linha_cabecalho: int) -> pd.DataFrame:
    return pd.read_excel(caminho, header=linha_cabecalho, dtype=str, engine='openpyxl')


def _normalizar_codigo(valor_raw: str) -> str | None:
    valor = str(valor_raw).strip()
    if valor.upper() in ("NAN", "NAT", "<NA>", "NONE", ""):
        return None
    if valor.endswith(".0"):
        valor = valor[:-2]
    if not valor.replace(".", "").isdigit():
        return None
    return valor.zfill(10)


def _eh_linha_de_dados(row: pd.Series, col_numero: str) -> bool:
    val = str(row.get(col_numero, "")).strip()
    if val.upper() in ("NAN", "NAT", "<NA>", "NONE", ""):
        return False
    if val.endswith(".0"):
        val = val[:-2]
    return val.isdigit()


def _converter_float(valor) -> float | None:
    try:
        return float(str(valor).strip().replace(",", "."))
    except (ValueError, TypeError):
        return None


def _validar_percentual(valor_base: float, valor_perc: float, percentual: float) -> str:
    esperado = round(valor_base * percentual, 2)
    if abs(valor_perc - esperado) < TOLERANCIA_CENTAVOS:
        return "OK"
    return f"Divergente (esperado: R$ {esperado:.2f}, encontrado: R$ {valor_perc:.2f})"


def _auditar_linha(row: pd.Series, cols: dict, dict_sigtap: dict) -> dict:
    cod_raw        = str(row.get(cols["codigo"], ""))
    cod_sigtap     = _normalizar_codigo(cod_raw)
    valor_proc     = _converter_float(row.get(cols["valor_proc"]))
    valor_principal = _converter_float(row.get(cols["valor_principal"]))
    valor_auxiliar  = _converter_float(row.get(cols["valor_auxiliar"]))

    resultado = {
        "Nº":                         str(row.get(cols["numero"], "")).replace(".0", ""),
        "Paciente":                   row.get(cols.get("paciente", ""), ""),
        "Cirurgião Principal":        row.get(cols.get("principal", ""), ""),
        "Médico Auxiliar":            row.get(cols.get("auxiliar", ""), ""),
        "Código SIGTAP":              cod_sigtap if cod_sigtap else cod_raw,
        "Descrição (Planilha)":       row.get(cols.get("descricao", ""), ""),
        "Valor Procedimento (Planilha)": valor_proc,
        "Valor Procedimento (SIGTAP)":   None,
        "Status Valor Procedimento":     "",
        "Diferença Valor (R$)":          None,
        "Valor Cirurgião (75%) - Planilha": valor_principal,
        "Status 75% Cirurgião":          "",
        "Valor Auxiliar (25%) - Planilha":  valor_auxiliar,
        "Status 25% Auxiliar":           "",
        "Status CBO":                    "",
    }

    if cod_sigtap is None:
        resultado.update({
            "Status CBO": "Código inválido",
            "Status Valor Procedimento": "N/A",
            "Status 75% Cirurgião": "N/A",
            "Status 25% Auxiliar": "N/A",
        })
        return resultado

    if cod_sigtap not in dict_sigtap:
        resultado.update({
            "Status CBO": "Procedimento NÃO previsto para o CBO",
            "Status Valor Procedimento": "N/A (CBO inválido)",
            "Status 75% Cirurgião": "N/A",
            "Status 25% Auxiliar": "N/A",
        })
        return resultado

    resultado["Status CBO"] = "Previsto para o CBO"
    valor_sigtap = float(dict_sigtap[cod_sigtap])
    resultado["Valor Procedimento (SIGTAP)"] = valor_sigtap

    if valor_proc is None:
        resultado["Status Valor Procedimento"] = "Valor ausente na planilha"
    else:
        diferenca = valor_proc - valor_sigtap
        resultado["Diferença Valor (R$)"] = diferenca
        resultado["Status Valor Procedimento"] = (
            "Valor correto" if abs(diferenca) < TOLERANCIA_CENTAVOS else "Valor DIVERGENTE"
        )

    resultado["Status 75% Cirurgião"] = (
        _validar_percentual(valor_proc, valor_principal, PERCENTUAL_PRINCIPAL)
        if valor_proc is not None and valor_principal is not None else "Valor ausente"
    )
    resultado["Status 25% Auxiliar"] = (
        _validar_percentual(valor_proc, valor_auxiliar, PERCENTUAL_AUXILIAR)
        if valor_proc is not None and valor_auxiliar is not None else "Valor ausente"
    )

    return resultado


def _salvar_relatorio(df_resultado: pd.DataFrame, arquivo_saida: str) -> None:
    df_resultado.to_excel(arquivo_saida, index=False)
    wb = load_workbook(arquivo_saida)
    ws = wb.active

    fmt_moeda  = 'R$ #,##0.00'
    fill_ok    = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    fill_erro  = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    fill_aviso = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

    cols_moeda  = ["Valor Procedimento (Planilha)", "Valor Procedimento (SIGTAP)",
                   "Diferença Valor (R$)", "Valor Cirurgião (75%) - Planilha", "Valor Auxiliar (25%) - Planilha"]
    cols_status = ["Status Valor Procedimento", "Status 75% Cirurgião", "Status 25% Auxiliar", "Status CBO"]

    header = [c.value for c in ws[1]]
    idx_moeda  = [header.index(c) + 1 for c in cols_moeda  if c in header]
    idx_status = [(c, header.index(c) + 1) for c in cols_status if c in header]

    for row_idx in range(2, len(df_resultado) + 2):
        for ci in idx_moeda:
            cell = ws.cell(row_idx, ci)
            if isinstance(cell.value, (int, float)):
                cell.number_format = fmt_moeda
        for col_name, ci in idx_status:
            cell = ws.cell(row_idx, ci)
            val  = str(cell.value or "").upper()
            if "CORRETO" in val or val == "OK" or ("PREVISTO" in val and "NÃO" not in val):
                cell.fill = fill_ok
            elif "DIVERGENTE" in val or "NÃO PREVISTO" in val or "INVÁLIDO" in val:
                cell.fill = fill_erro
                cell.font = Font(bold=True)
            elif "AUSENTE" in val or "N/A" in val:
                cell.fill = fill_aviso

    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    wb.save(arquivo_saida)


# =====================================================================
# Funcao principal programatica
# =====================================================================

def executar(planilha_path: str, cbo: str = CBO_NEUROCIRURGIAO, mapeamento: dict = None, linha_cabecalho: int = None) -> dict:
    """
    Executa a auditoria Modelo 02 de forma programatica.

    Args:
        planilha_path : Caminho completo para a planilha no formato PLENA.
        cbo           : Codigo CBO alvo.
        mapeamento    : Dicionario com os nomes reais das colunas na planilha.
        linha_cabecalho: Indice da linha de cabecalho. Se None, detecta automaticamente.
    """
    if not os.path.exists(planilha_path):
        return {"sucesso": False, "erro": f"Planilha '{planilha_path}' nao encontrada."}

    try:
        dict_sigtap = _carregar_base_sigtap(cbo)
    except FileNotFoundError as e:
        return {"sucesso": False, "erro": str(e)}

    agora = datetime.now()
    pasta_auditoria = os.path.join("auditoria", agora.strftime("%Y-%m-%d"))
    os.makedirs(pasta_auditoria, exist_ok=True)
    arquivo_saida = os.path.join(
        pasta_auditoria,
        f"auditoria_modelo02_CBO_{cbo}_{agora.strftime('%H-%M-%S')}.xlsx"
    )

    if linha_cabecalho is None:
        linha_cabecalho = _detectar_cabecalho(planilha_path)
        
    df_plena  = _carregar_planilha_plena(planilha_path, linha_cabecalho)
    col_list  = df_plena.columns.tolist()

    # --- Mapeamento de colunas ---
    # Prioridade: mapeamento manual > auto-deteccao por nome > posicao
    if mapeamento:
        cols = {
            "numero":          mapeamento.get("numero"),
            "codigo":          mapeamento.get("codigo"),
            "valor_proc":      mapeamento.get("valor_proc"),
            "valor_principal": mapeamento.get("valor_principal"),
            "valor_auxiliar":  mapeamento.get("valor_auxiliar"),
            # Campos opcionais: auto-deteccao por nome ou posicao
            "paciente":   next((c for c in col_list if "PACIENTE" in str(c).upper()), col_list[1] if len(col_list) > 1 else None),
            "principal":  next((c for c in col_list if "CIRURGI" in str(c).upper() and "PRINCIPAL" in str(c).upper() and "PAGAR" not in str(c).upper()), col_list[5] if len(col_list) > 5 else None),
            "auxiliar":   next((c for c in col_list if "AUXILIAR" in str(c).upper() and "PAGAR" not in str(c).upper()), col_list[6] if len(col_list) > 6 else None),
            "descricao":  next((c for c in col_list if "DESCRI" in str(c).upper() and "PROCEDIMENTO" in str(c).upper()), col_list[8] if len(col_list) > 8 else None),
        }
    else:
        # Auto-deteccao completa por nome
        cols = {}
        for col in col_list:
            cu = str(col).upper()
            if cu.strip() in ("N°", "N\uFFFD", "NO", "N"):
                cols.setdefault("numero", col)
            elif "PACIENTE" in cu:
                cols.setdefault("paciente", col)
            elif "CIRURGI" in cu and "PRINCIPAL" in cu and "PAGAR" in cu:
                cols.setdefault("valor_principal", col)
            elif "AUXILIAR" in cu and "PAGAR" in cu:
                cols.setdefault("valor_auxiliar", col)
            elif "CIRURGI" in cu and "PRINCIPAL" in cu:
                cols.setdefault("principal", col)
            elif "AUXILIAR" in cu:
                cols.setdefault("auxiliar", col)
            elif ("COD" in cu or "CÓD" in cu) and "SIGTAP" in cu:
                cols.setdefault("codigo", col)
            elif "DESCRI" in cu and "PROCEDIMENTO" in cu:
                cols.setdefault("descricao", col)
            elif "VALOR DO PROCEDIMENTO" in cu:
                cols.setdefault("valor_proc", col)

        # Fallback por posicao
        defaults = {"numero": 0, "paciente": 1, "principal": 5, "auxiliar": 6,
                    "codigo": 7, "descricao": 8, "valor_proc": 9,
                    "valor_principal": 10, "valor_auxiliar": 11}
        for key, idx in defaults.items():
            if key not in cols and len(col_list) > idx:
                cols[key] = col_list[idx]

    # Verificar colunas obrigatorias
    obrigatorias = ["numero", "codigo", "valor_proc", "valor_principal", "valor_auxiliar"]
    faltando = [k for k in obrigatorias if not cols.get(k)]
    if faltando:
        disponiveis = [str(c) for c in col_list if not str(c).startswith("Unnamed")]
        return {
            "sucesso": False,
            "erro": (
                f"Colunas obrigatorias nao mapeadas: {faltando}. "
                f"Colunas disponiveis: {disponiveis}"
            )
        }

    resultados = []
    linhas_ignoradas = 0

    for _, row in df_plena.iterrows():
        if not _eh_linha_de_dados(row, cols["numero"]):
            linhas_ignoradas += 1
            continue
        resultados.append(_auditar_linha(row, cols, dict_sigtap))

    if not resultados:
        return {"sucesso": False, "erro": "Nenhum procedimento encontrado. Verifique se as colunas estao mapeadas corretamente."}

    df_resultado = pd.DataFrame(resultados)
    _salvar_relatorio(df_resultado, arquivo_saida)

    total         = len(df_resultado)
    corretos      = int((df_resultado["Status Valor Procedimento"] == "Valor correto").sum())
    divergentes   = int(df_resultado["Status Valor Procedimento"].str.contains("DIVERGENTE", na=False).sum())
    nao_previstos = int(df_resultado["Status CBO"].str.contains("NÃO previsto", case=False, na=False).sum())
    perc_ok       = int((df_resultado["Status 75% Cirurgião"] == "OK").sum())
    aux_ok        = int((df_resultado["Status 25% Auxiliar"] == "OK").sum())

    return {
        "sucesso": True,
        "total": total,
        "corretos": corretos,
        "divergentes": divergentes,
        "nao_previstos": nao_previstos,
        "perc_principal_ok": perc_ok,
        "perc_auxiliar_ok": aux_ok,
        "arquivo_saida": os.path.abspath(arquivo_saida),
    }


# =====================================================================
# Entrada CLI
# =====================================================================

def main():
    parser = argparse.ArgumentParser(description="Auditoria de valores SIGTAP - Modelo 02 (Planilha PLENA)")
    parser.add_argument("--planilha",           default="docs/NEURO JULHO_PLENA.xlsx")
    parser.add_argument("--cbo",                default=CBO_NEUROCIRURGIAO)
    parser.add_argument("--col_numero",         default=None)
    parser.add_argument("--col_codigo",         default=None)
    parser.add_argument("--col_valor_proc",     default=None)
    parser.add_argument("--col_valor_principal",default=None)
    parser.add_argument("--col_valor_auxiliar", default=None)
    args = parser.parse_args()

    mapeamento = None
    if all([args.col_numero, args.col_codigo, args.col_valor_proc,
            args.col_valor_principal, args.col_valor_auxiliar]):
        mapeamento = {
            "numero":          args.col_numero,
            "codigo":          args.col_codigo,
            "valor_proc":      args.col_valor_proc,
            "valor_principal": args.col_valor_principal,
            "valor_auxiliar":  args.col_valor_auxiliar,
        }

    resultado = executar(args.planilha, args.cbo, mapeamento)

    if not resultado["sucesso"]:
        print(f"[!] Erro: {resultado['erro']}")
        return

    print("\n" + "=" * 65)
    print("  AUDITORIA MODELO 02 — RESUMO")
    print("=" * 65)
    print(f"  Total auditados          : {resultado['total']}")
    print(f"  Valores corretos (Col J) : {resultado['corretos']}")
    print(f"  Valores DIVERGENTES      : {resultado['divergentes']}")
    print(f"  Fora do CBO              : {resultado['nao_previstos']}")
    print(f"  75% Cirurgiao OK         : {resultado['perc_principal_ok']}")
    print(f"  25% Auxiliar OK          : {resultado['perc_auxiliar_ok']}")
    print(f"\n  Relatorio: {resultado['arquivo_saida']}")
    print("=" * 65)


if __name__ == "__main__":
    main()
