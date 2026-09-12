"""
Busca os valores financeiros (SA, SH, SP) dos procedimentos.

Le uma planilha existente, faz requisicoes para detalhar cada procedimento
e extrai os valores, gerando uma nova planilha com essas colunas adicionais
formatadas como moeda.

Uso:
    python -m src.buscar_valores_procedimentos
"""
import argparse
import time
import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import numbers

from src.config import PAUSA_ENTRE_PAGINAS, CBO_NEUROCIRURGIAO, PASTA_DADOS
from src.sigtap_client import detalhar_procedimento
from src.xml_parser import extrair_valores


def main():
    parser = argparse.ArgumentParser(description="Busca valores financeiros na API SIGTAP")
    parser.add_argument("--cbo", default=CBO_NEUROCIRURGIAO,
                        help="CBO alvo (define os nomes dos arquivos automaticamente)")
    parser.add_argument("--coluna", default="CODIGO",
                        help="Nome da coluna que contem os codigos")
    args = parser.parse_args()

    entrada = os.path.join(PASTA_DADOS, f"exclusivo - procedimentos_cbo_{args.cbo}.xlsx")
    saida = os.path.join(PASTA_DADOS, f"exclusivo - procedimentos_cbo_{args.cbo}_com_valores.xlsx")

    print("=" * 60, flush=True)
    print(f"  Busca de Valores Financeiros SIGTAP - CBO {args.cbo}", flush=True)
    print("=" * 60 + "\n", flush=True)

    if not os.path.exists(entrada):
        print(f"[!] Erro: Arquivo '{entrada}' nao encontrado.", flush=True)
        return

    print(f"Lendo planilha: {entrada}", flush=True)
    # Ler forçando a coluna a ser string para não perder os zeros à esquerda,
    # porém caso ela já tenha sido lida sem zero pelo usuário antes, garantimos com zfill(10)
    df = pd.read_excel(entrada, dtype={args.coluna: str})
    
    # 1. Garante que os zeros à esquerda permaneçam na base de dados
    df[args.coluna] = df[args.coluna].astype(str).str.strip().str.zfill(10)

    total = len(df)
    print(f"Total de procedimentos para consultar: {total}\n", flush=True)

    # Estimativa de tempo
    tempo_estimado_seg = total * (1.5 + PAUSA_ENTRE_PAGINAS)
    tempo_estimado_min = tempo_estimado_seg / 60.0
    print(f"Tempo estimado de execucao: ~{tempo_estimado_min:.1f} minutos.\n", flush=True)

    valores_sa = []
    valores_sh = []
    valores_sp = []

    for idx, row in df.iterrows():
        codigo = row[args.coluna]
        
        # Para ser mais rapido, detalhamos a DESCRICAO
        xml_resp = detalhar_procedimento(codigo, categoria_detalhe="DESCRICAO", quantidade_cbos=1)

        valores = extrair_valores(xml_resp)
        valores_sa.append(valores.get("valorSA", 0.0))
        valores_sh.append(valores.get("valorSH", 0.0))
        valores_sp.append(valores.get("valorSP", 0.0))

        print(f"  [{idx + 1}/{total}] {codigo} | SA: {valores['valorSA']:.2f} | SH: {valores['valorSH']:.2f} | SP: {valores['valorSP']:.2f}", flush=True)

        time.sleep(PAUSA_ENTRE_PAGINAS)

    df["valorSA"] = valores_sa
    df["valorSH"] = valores_sh
    df["valorSP"] = valores_sp

    # Salva inicialmente com o Pandas
    df.to_excel(saida, index=False)
    
    # 2. Formata como moeda (Contábil/Financeiro) usando openpyxl
    wb = load_workbook(saida)
    ws = wb.active
    
    # Descobre o índice das colunas de valores inseridas no final
    col_idx_sa = df.columns.get_loc("valorSA") + 1
    col_idx_sh = df.columns.get_loc("valorSH") + 1
    col_idx_sp = df.columns.get_loc("valorSP") + 1
    
    formato_moeda = 'R$ #,##0.00'
    
    # Aplica formatação de moeda em todas as linhas das colunas de valores
    for row_idx in range(2, len(df) + 2):  # Pula o cabeçalho
        ws.cell(row=row_idx, column=col_idx_sa).number_format = formato_moeda
        ws.cell(row=row_idx, column=col_idx_sh).number_format = formato_moeda
        ws.cell(row=row_idx, column=col_idx_sp).number_format = formato_moeda

    wb.save(saida)

    print(f"\n{'=' * 60}", flush=True)
    print(f"Sucesso! Resultado (com zeros à esquerda e formato moeda) salvo em '{saida}'.", flush=True)


if __name__ == "__main__":
    main()
